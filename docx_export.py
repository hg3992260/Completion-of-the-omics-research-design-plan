# -*- coding: utf-8 -*-
"""把研究设计导出成 Word（.docx）：带标题层级、检查表用表格、问答用两列表。

用 python-docx 纯 Python 实现，不依赖 Office / officecli，因此能一起打包进 exe/app。
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK_FONT = "微软雅黑"          # Windows 有；macOS 的 Word 会自动替换成中文字体
ASCII_FONT = "Calibri"
STATUS_MARK = {"done": "✅ 已完成", "drafted": "◐ 待采纳", "asked": "◔ 已追问", "todo": "○ 未开始"}


def _set_font(style, size=None, bold=None, color=None):
    style.font.name = ASCII_FONT
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)


def _init_styles(doc: Document) -> None:
    """设定中英文字体，避免中文在 Word 里变宋体/方框。"""
    _set_font(doc.styles["Normal"], size=10.5)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13),
                       ("Heading 3", 11.5), ("Heading 4", 10.5)):
        try:
            _set_font(doc.styles[name], size=size, bold=True)
        except KeyError:
            pass
    for name in ("Caption", "Quote", "Intense Quote"):
        try:
            _set_font(doc.styles[name], size=9)
        except KeyError:
            pass


def _para(doc, text: str, style=None, size=None, italic=False, color=None, space_after=6):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _body(doc, text: str) -> None:
    """把多行正文拆成段落；以 - / • 开头的当作项目符号。"""
    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.lstrip().startswith(("- ", "• ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line.lstrip()[2:].strip())
            rpr = run._element.get_or_add_rPr()
            rpr.get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
        else:
            _para(doc, line, space_after=4)


def _table(doc, headers: list, rows: list) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(project, path: str) -> str:
    """把 Project 渲染成 .docx 并写盘，返回最终路径。"""
    from design_agent import q_text
    from stages_data import STAGES

    doc = Document()
    _init_styles(doc)

    doc.add_paragraph(project.name, style="Title")
    _para(doc, f"由「组学研究设计工作台」生成　·　{project.created}　·　模型 "
               f"{project.model or '—'}",
          size=9, italic=True, color=(0x60, 0x60, 0x60), space_after=14)

    # ---------------- 原始输入
    doc.add_heading("一、研究设想（原始输入）", level=1)
    if (project.raw_design or "").strip():
        _body(doc, project.raw_design)
    else:
        _para(doc, "（未填写）", italic=True, color=(0x88, 0x88, 0x88))

    # ---------------- 分阶段
    doc.add_heading("二、分阶段完善稿", level=1)
    counts = project.status_counts()
    _para(doc, f"完成度：已定稿 {counts['done']}/10　待采纳 {counts['drafted']}　"
               f"已追问 {counts['asked']}　未开始 {counts['todo']}",
          size=9, color=(0x60, 0x60, 0x60), space_after=10)

    for s in STAGES:
        st = project.stage(s["id"])
        status = st.get("status", "todo")
        body = (st.get("final") or st.get("draft") or "").strip()

        h = doc.add_heading(level=2)
        run = h.add_run(f"{s['id']:02d}　{s['title']}")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
        run2 = h.add_run(f"　[{STATUS_MARK.get(status, status)}]")
        run2.font.size = Pt(9)
        run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        run2._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
        _para(doc, f"规范出处：{s['spec']}", size=8.5, italic=True,
              color=(0x70, 0x70, 0x70), space_after=6)

        if body:
            _body(doc, body)
        else:
            _para(doc, "（尚未完成本阶段）", italic=True, color=(0x88, 0x88, 0x88))

        questions = st.get("questions") or []
        answers = st.get("answers") or []
        if questions:
            doc.add_heading("追问与回答", level=3)
            rows = []
            for i, q in enumerate(questions):
                ans = answers[i] if i < len(answers) else ""
                rows.append((q_text(q), ans or "（未回答）"))
            _table(doc, ["问题", "回答"], rows)

        checks = st.get("checklist") or []
        if checks:
            doc.add_heading("检查表", level=3)
            rows = []
            for i, c in enumerate(checks, 1):
                text = c if isinstance(c, str) else str(c)
                if "｜" in text:
                    item, ref = text.split("｜", 1)
                    rows.append((i, item.strip(), ref.strip()))
                else:
                    rows.append((i, text.strip(), ""))
            _table(doc, ["#", "条目", "依据"], rows)

        risks = (st.get("risks") or "").strip()
        if risks:
            doc.add_heading("风险提示", level=3)
            _para(doc, risks, color=(0xA0, 0x30, 0x30))

    # ---------------- 完整草案
    if (getattr(project, "final_doc", "") or "").strip():
        doc.add_page_break()
        doc.add_heading("三、完整研究设计草案", level=1)
        _body(doc, project.final_doc)

    if not path.lower().endswith(".docx"):
        path += ".docx"
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":                                        # 手动试跑
    import sys

    from design_agent import Project

    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join("projects", "未命名课题.json")
    out = sys.argv[2] if len(sys.argv) > 2 else "研究设计_试跑.docx"
    print(build(Project.load(src), out))
