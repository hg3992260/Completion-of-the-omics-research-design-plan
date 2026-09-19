# -*- coding: utf-8 -*-
"""检查导出的 docx 结构：标题层级、表格、中文与内容完整性。"""
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
path = sys.argv[1] if len(sys.argv) > 1 else "_试跑.docx"
doc = Document(path)

heads = {}
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        heads[p.style.name] = heads.get(p.style.name, 0) + 1

print(f"文件: {path}")
print(f"段落总数: {len(doc.paragraphs)}　表格数: {len(doc.tables)}")
print("标题层级统计:", dict(sorted(heads.items())))
print("\n前 8 个非空段落:")
shown = 0
for p in doc.paragraphs:
    if p.text.strip():
        print(f"   [{p.style.name}] {p.text.strip()[:70]}")
        shown += 1
        if shown >= 8:
            break

print("\n表格概况:")
for i, t in enumerate(doc.tables, 1):
    head = " | ".join(c.text.strip()[:14] for c in t.rows[0].cells)
    print(f"   表{i}: {len(t.rows)} 行 × {len(t.columns)} 列　表头：{head}")
    if len(t.rows) > 1:
        print(f"        首行数据：{' | '.join(c.text.strip()[:30] for c in t.rows[1].cells)}")

full = "\n".join(p.text for p in doc.paragraphs)
for key in ("研究设想（原始输入）", "分阶段完善稿", "追问与回答", "检查表", "规范出处", "完成度"):
    print(f"含「{key}」: {key in full}")
# 检查中文是否正常（不是方框/乱码）
cjk = sum(1 for ch in full if "\u4e00" <= ch <= "\u9fff")
print(f"中文字符数: {cjk}")
