# -*- coding: utf-8 -*-
"""实测「导出 Word」按钮：对话框默认值、写盘、docx 结构、取消、MCP 工具。"""
import os
import sys

sys.argv = ["x"]
from PySide6.QtWidgets import QApplication, QFileDialog
from PySide6 import QtCore

import design_agent as da
import design_studio as ds

app = QApplication([])
ds.set_color_theme(ds.THEME_PATH)
ds.set_appearance_mode("light")
win = ds.StudioWindow(demo=False)
win.resize(1638, 830)
p = os.path.join(da.PROJECT_DIR, "未命名课题.json")
if os.path.exists(p):
    win.project = da.Project.load(p)
    win._refresh_all()
win.show()

out = []
TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_docx_tmp")
os.makedirs(TMP, exist_ok=True)


def stub(return_path):
    box = {}

    def fake(parent, caption, default, filters, *a, **kw):
        box.update(caption=caption, default=default, filters=filters)
        return (return_path, "")

    QFileDialog.getSaveFileName = staticmethod(fake)
    return box


def run():
    # 1) 正常导出
    target = os.path.join(TMP, "伦理申请附件.docx")
    box = stub(target)
    win.export_docx()
    out.append("【1】导出 Word")
    out.append(f"  对话框标题：{box.get('caption')!r}")
    out.append(f"  默认文件名：{os.path.basename(box.get('default',''))}")
    out.append(f"  过滤器：{box.get('filters')!r}")
    ok = os.path.exists(target)
    out.append(f"  文件已生成：{ok}　大小 {os.path.getsize(target) if ok else 0} 字节")
    out.append(f"  页脚提示：{win.tip.label().text()!r}")
    if ok:
        from docx import Document
        d = Document(target)
        heads = {}
        for para in d.paragraphs:
            if para.style.name.startswith("Heading"):
                heads[para.style.name] = heads.get(para.style.name, 0) + 1
        out.append(f"  结构：段落 {len(d.paragraphs)}　表格 {len(d.tables)}　标题 {dict(sorted(heads.items()))}")
        if d.tables:
            t = d.tables[0]
            out.append(f"  首个表格：{len(t.rows)}×{len(t.columns)}　表头 "
                       f"{' | '.join(c.text for c in t.rows[0].cells)}")

    # 2) 取消
    stub("")
    before = set(os.listdir(TMP))
    win.export_docx()
    out.append("【2】取消")
    out.append(f"  是否新增文件：{bool(set(os.listdir(TMP)) - before)}")
    out.append(f"  页脚提示：{win.tip.label().text()!r}")

    # 3) 不写扩展名
    stub(os.path.join(TMP, "无扩展名"))
    win.export_docx()
    out.append("【3】未写扩展名 → 自动补 .docx："
               f"{os.path.exists(os.path.join(TMP, '无扩展名.docx'))}")

    # 4) MCP 工具
    from mcp_server import export_docx
    import json as _json
    r = _json.loads(export_docx(p.replace(".json", "")))
    out.append("【4】MCP 工具 export_docx")
    out.append(f"  {r}")

    open("_docx_ui_test.txt", "w", encoding="utf-8").write("\n".join(out))
    print("\n".join(out))
    app.quit()


QtCore.QTimer.singleShot(1200, run)
sys.exit(app.exec())
